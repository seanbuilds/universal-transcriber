"""Atomic Multi-Format Transcript Exporter with ISO-8601 Naming and DOCX/WebVTT Support (v4).
<!-- v4 – Standardized ISO-8601 YYYYMMDD_<Title> naming, custom title overrides, W3C WebVTT, and atomic multi-format exports -->
"""

from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import List, Dict, Any, Optional

try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    docx = None

from config_v4 import TRANSCRIPTS_DIR


def format_srt_timestamp(seconds: float) -> str:
    """Format seconds into SRT timestamp format: HH:MM:SS,mmm."""
    seconds = max(0.0, float(seconds))
    hrs = int(seconds // 3600)
    mins = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    msec = int((seconds - int(seconds)) * 1000)
    return f"{hrs:02d}:{mins:02d}:{secs:02d},{msec:03d}"


def format_vtt_timestamp(seconds: float) -> str:
    """Format seconds into WebVTT timestamp format: HH:MM:SS.mmm."""
    seconds = max(0.0, float(seconds))
    hrs = int(seconds // 3600)
    mins = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    msec = int((seconds - int(seconds)) * 1000)
    return f"{hrs:02d}:{mins:02d}:{secs:02d}.{msec:03d}"


def sanitize_name(name: str) -> str:
    """Strictly sanitize string for safe folder and file naming (no traversal sequences or illegal chars)."""
    clean = re.sub(r'[^a-zA-Z0-9_\- ]', '', name).strip('. ')
    return clean[:60] or "Untitled"


def format_iso_date(raw_date: Optional[str] = None) -> str:
    """Extract or generate a standardized ISO-8601 date string (YYYYMMDD).
    
    If raw_date contains a recognizable date (e.g., '2026-09-17', '20260917', '2026.09.17'),
    standardizes to YYYYMMDD. Otherwise, defaults to current UTC date YYYYMMDD.
    """
    if raw_date:
        # Match YYYYMMDD, YYYY-MM-DD, YYYY_MM_DD
        m = re.search(r'(\d{4})[-_.]?(\d{2})[-_.]?(\d{2})', str(raw_date))
        if m:
            year, month, day = m.group(1), m.group(2), m.group(3)
            return f"{year}{month}{day}"

    return datetime.now(timezone.utc).strftime("%Y%m%d")


def build_iso_filename(
    custom_name: Optional[str] = None,
    metadata_title: Optional[str] = None,
    metadata_date: Optional[str] = None,
) -> str:
    """Build a standardized ISO filename in format: YYYYMMDD_<Title>.
    
    Example output: 20260917_Gaming_Highlight
    """
    iso_date = format_iso_date(metadata_date)

    chosen_title = ""
    if custom_name and custom_name.strip():
        chosen_title = custom_name.strip()
    elif metadata_title and metadata_title.strip():
        chosen_title = metadata_title.strip()
    else:
        chosen_title = "Recording"

    # If title already begins with an 8-digit date string, preserve that date
    m = re.match(r'^(\d{8})[_\- ]?(.*)$', chosen_title)
    if m:
        iso_date = m.group(1)
        sub_title = m.group(2).strip()
        chosen_title = sub_title if sub_title else "Recording"

    clean_title = sanitize_name(chosen_title).replace(" ", "_")
    return f"{iso_date}_{clean_title}"


def atomic_write_text(target_path: Path, content: str) -> None:
    """Write text atomically via temporary file and atomic rename."""
    target_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = target_path.with_suffix(target_path.suffix + ".tmp")
    temp_path.write_text(content, encoding="utf-8")
    temp_path.replace(target_path)


class TranscriptExporterV4:
    """Exports structured transcripts to Markdown, Plaintext, SRT, WebVTT, DOCX, and JSON using ISO naming."""

    def __init__(self, base_output_dir: Path = TRANSCRIPTS_DIR):
        self.base_output_dir = Path(base_output_dir)

    def export(
        self,
        blocks: List[Dict[str, Any]],
        metadata: Dict[str, Any],
        custom_dir: Optional[Path] = None,
        custom_name: Optional[str] = None,
        direct_dir: bool = False,
    ) -> Dict[str, Path]:
        """Save all six transcript formats to target folder atomically with ISO-8601 naming strictly after all pipeline steps succeed."""
        raw_title = metadata.get("title", "Untitled Recording")
        raw_date = metadata.get("date") or metadata.get("upload_date") or ""

        folder_name = build_iso_filename(
            custom_name=custom_name or metadata.get("custom_title"),
            metadata_title=raw_title,
            metadata_date=raw_date,
        )

        if direct_dir and custom_dir:
            target_dir = Path(custom_dir)
        else:
            target_dir = (custom_dir or self.base_output_dir) / folder_name

        dir_created = False
        if not target_dir.exists():
            target_dir.mkdir(parents=True, exist_ok=True)
            dir_created = True

        md_path = target_dir / f"{folder_name}.md"
        txt_path = target_dir / f"{folder_name}.txt"
        srt_path = target_dir / f"{folder_name}.srt"
        vtt_path = target_dir / f"{folder_name}.vtt"
        docx_path = target_dir / f"{folder_name}.docx"
        json_path = target_dir / f"{folder_name}.json"

        participants = sorted(list(set(b.get("speaker", "Speaker") for b in blocks if b.get("speaker"))))
        source_url = metadata.get("source") or metadata.get("url") or "N/A"

        try:
            # 1. Markdown
            md_lines = [
                f"# {raw_title}",
                "",
                f"**Meeting / Transcript Name**: {raw_title}  ",
                f"**YouTube / Source URL**: {source_url}  ",
                f"**Date**: {raw_date or 'N/A'}  ",
                f"**ISO Identifier**: {folder_name}  ",
                f"**Duration**: {metadata.get('duration_str', 'N/A')}  ",
                "",
                "### Participants",
                "",
            ]
            for p in participants:
                md_lines.append(f"- **{p}**")
            md_lines.extend([
                "",
                "---",
                "",
                "## Verbatim Transcript",
                "",
            ])
            for b in blocks:
                md_lines.append(f"### [{b.get('ts', '00:00:00')}] {b.get('speaker', 'Speaker')}")
                md_lines.append("")
                md_lines.append(b.get("text", ""))
                md_lines.append("")

            atomic_write_text(md_path, "\n".join(md_lines))

            # 2. Plaintext
            txt_lines = [
                "=" * 80,
                f"Meeting / Transcript Name: {raw_title}",
                f"YouTube / Source URL: {source_url}",
                f"ISO Identifier: {folder_name}",
                f"Date: {raw_date or 'N/A'}",
                f"Duration: {metadata.get('duration_str', 'N/A')}",
                "Participants:",
            ]
            for p in participants:
                txt_lines.append(f"  - {p}")
            txt_lines.extend([
                "=" * 80,
                "",
            ])
            for b in blocks:
                txt_lines.append(f"[{b.get('ts', '00:00:00')}] {b.get('speaker', 'Speaker')}: {b.get('text', '')}")
                txt_lines.append("")

            atomic_write_text(txt_path, "\n".join(txt_lines))

            # 3. Subtitles (.srt)
            srt_lines = []
            for idx, b in enumerate(blocks, 1):
                s_val = max(0.0, float(b.get("start", 0.0)))
                e_val = max(s_val + 0.5, float(b.get("end", s_val + 3.0)))
                s_time = format_srt_timestamp(s_val)
                e_time = format_srt_timestamp(e_val)
                srt_lines.append(str(idx))
                srt_lines.append(f"{s_time} --> {e_time}")
                srt_lines.append(f"{b.get('speaker', 'Speaker')}: {b.get('text', '')}")
                srt_lines.append("")

            atomic_write_text(srt_path, "\n".join(srt_lines))

            # 4. WebVTT (.vtt) with W3C Voice Span Cues
            vtt_lines = [
                "WEBVTT - Universal Transcriber",
                "",
                f"NOTE Meeting / Transcript Name: {raw_title}",
                f"NOTE YouTube / Source URL: {source_url}",
                f"NOTE Identifier: {folder_name}",
                f"NOTE Date: {raw_date or 'N/A'}",
                "",
            ]
            for idx, b in enumerate(blocks, 1):
                s_val = max(0.0, float(b.get("start", 0.0)))
                e_val = max(s_val + 0.5, float(b.get("end", s_val + 3.0)))
                s_time = format_vtt_timestamp(s_val)
                e_time = format_vtt_timestamp(e_val)
                spk = b.get("speaker", "Speaker")
                text_payload = b.get("text", "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                vtt_lines.append(f"{idx}")
                vtt_lines.append(f"{s_time} --> {e_time}")
                vtt_lines.append(f"<v {spk}>{text_payload}</v>")
                vtt_lines.append("")

            atomic_write_text(vtt_path, "\n".join(vtt_lines))

            # 5. Native DOCX Document
            if docx is not None:
                doc = docx.Document()
                title_heading = doc.add_heading(raw_title, level=1)
                title_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

                meta_p = doc.add_paragraph()
                meta_p.add_run(f"Meeting / Transcript Name: {raw_title}\n")
                meta_p.add_run(f"YouTube / Source URL: {source_url}\n")
                meta_p.add_run(f"ISO Identifier: {folder_name}\n")
                meta_p.add_run(f"Date: {raw_date or 'N/A'}\n")
                meta_p.add_run(f"Duration: {metadata.get('duration_str', 'N/A')}\n")
                meta_p.add_run(f"Total Turns: {len(blocks)}\n")

                if participants:
                    doc.add_heading("Participants", level=2)
                    for p in participants:
                        doc.add_paragraph(f"• {p}", style="List Bullet" if "List Bullet" in doc.styles else None)

                doc.add_heading("Verbatim Transcript", level=2)
                for b in blocks:
                    tp = doc.add_paragraph()
                    run_hdr = tp.add_run(f"[{b.get('ts', '00:00:00')}] {b.get('speaker', 'Speaker')}\n")
                    run_hdr.bold = True
                    tp.add_run(b.get("text", ""))

                temp_docx = docx_path.with_suffix(".docx.tmp")
                doc.save(str(temp_docx))
                temp_docx.replace(docx_path)
            else:
                atomic_write_text(docx_path, "python-docx not installed.")

            # 6. JSON
            json_data = {
                "meeting_name": raw_title,
                "source_url": source_url,
                "iso_name": folder_name,
                "metadata": metadata,
                "participants": participants,
                "total_blocks": len(blocks),
                "transcript": blocks,
            }
            atomic_write_text(json_path, json.dumps(json_data, indent=2))

            # Verify all six offered formats exist locally on disk
            result_files = {
                "dir": target_dir,
                "iso_name": folder_name,
                "md": md_path,
                "txt": txt_path,
                "srt": srt_path,
                "vtt": vtt_path,
                "docx": docx_path,
                "json": json_path,
            }
            for fmt in ["md", "txt", "srt", "vtt", "docx", "json"]:
                fpath = result_files[fmt]
                if not fpath.exists():
                    raise RuntimeError(f"Failed to verify local {fmt.upper()} export file at: {fpath}")

            return result_files

        except Exception:
            # Atomic cleanup: if export fails, remove any created files and clean empty target dir
            for p in [md_path, txt_path, srt_path, vtt_path, docx_path, json_path]:
                if p.exists():
                    try:
                        p.unlink()
                    except OSError:
                        pass
                tmp_p = p.with_suffix(p.suffix + ".tmp")
                if tmp_p.exists():
                    try:
                        tmp_p.unlink()
                    except OSError:
                        pass
            if dir_created and target_dir.exists():
                try:
                    if not any(target_dir.iterdir()):
                        target_dir.rmdir()
                except OSError:
                    pass
            raise


TranscriptExporter = TranscriptExporterV4

"""Tests for Phase 4: Batch Catalog Ingestion, Archival, DOCX & WebVTT Exporters (v4).
<!-- v4 – Verification of DOCX, WebVTT, RSS parsing, and SQLite catalog idempotency -->
"""

import json
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

try:
    import docx
except ImportError:
    docx = None

from src.engine.export_v3 import TranscriptExporterV3, format_vtt_timestamp
from src.engine.catalog_v1 import MediaCatalog
from src.engine.ingest_v3 import MediaIngestorV3
try:
    from src.engine.pipeline_v4 import TranscriptionPipelineV4
except ImportError:
    from src.engine.archive.pipeline_v4 import TranscriptionPipelineV4


class TestPhase4CatalogExport(unittest.TestCase):

    def setUp(self):
        self.tmp_dir = Path(tempfile.mkdtemp())
        self.catalog_db = self.tmp_dir / "catalog_test.sqlite"
        self.catalog = MediaCatalog(db_path=self.catalog_db)
        self.ingestor = MediaIngestorV3(temp_dir=self.tmp_dir, catalog=self.catalog)
        self.exporter = TranscriptExporterV3(base_output_dir=self.tmp_dir)

    def test_format_vtt_timestamp_spec_compliance(self):
        """WebVTT timestamps require period as millisecond separator: HH:MM:SS.mmm."""
        ts = format_vtt_timestamp(65.123)
        self.assertEqual(ts, "00:01:05.123")
        self.assertIn(".", ts)
        self.assertNotIn(",", ts)

    def test_export_webvtt_has_voice_spans(self):
        """Exported .vtt file must contain WEBVTT header and voice span cues <v Speaker>."""
        blocks = [
            {"ts": "00:00:00", "start": 0.0, "end": 3.5, "speaker": "Chair Corey Evans", "text": "Good evening."},
            {"ts": "00:00:04", "start": 4.0, "end": 6.2, "speaker": "Craig MacLellan", "text": "Present."}
        ]
        meta = {"title": "Town Meeting", "date": "2026-09-16", "source": "local"}
        paths = self.exporter.export(blocks, meta)

        vtt_file = paths["vtt"]
        self.assertTrue(vtt_file.exists())
        content = vtt_file.read_text(encoding="utf-8")

        self.assertTrue(content.startswith("WEBVTT"))
        self.assertIn("<v Chair Corey Evans>Good evening.</v>", content)
        self.assertIn("<v Craig MacLellan>Present.</v>", content)
        self.assertIn("00:00:00.000 --> 00:00:03.500", content)

    def test_export_docx_document_structure(self):
        """Exported .docx file must contain formatted headings, metadata, and speaker blocks."""
        if docx is None:
            self.skipTest("python-docx not installed")

        blocks = [
            {"ts": "00:01:10", "start": 70.0, "end": 75.0, "speaker": "Chair", "text": "Motion is approved."},
        ]
        meta = {"title": "Executive Session", "date": "2026-09-16", "source": "local", "duration_str": "00:01:15"}
        paths = self.exporter.export(blocks, meta)

        docx_file = paths["docx"]
        self.assertTrue(docx_file.exists())

        doc = docx.Document(str(docx_file))
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        all_text = "\n".join(p.text for p in doc.paragraphs)

        self.assertIn("Executive Session", headings)
        self.assertIn("Verbatim Transcript", headings)
        self.assertIn("Date: 2026-09-16", all_text)
        self.assertIn("[00:01:10] Chair", all_text)
        self.assertIn("Motion is approved.", all_text)

    def test_rss_podcast_feed_parsing(self):
        """RSS parser extracts items with audio enclosures and titles."""
        sample_rss = """<?xml version="1.0" encoding="UTF-8"?>
        <rss version="2.0">
          <channel>
            <title>Select Board Audio</title>
            <item>
              <title>Regular Meeting 09/16</title>
              <guid>meeting_0916</guid>
              <pubDate>Wed, 16 Sep 2026 18:00:00 GMT</pubDate>
              <enclosure url="https://example.com/audio_0916.mp3" type="audio/mpeg" length="45000000"/>
            </item>
          </channel>
        </rss>
        """
        items = self.ingestor.parse_rss_feed(sample_rss)
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["id"], "meeting_0916")
        self.assertEqual(items[0]["title"], "Regular Meeting 09/16")
        self.assertEqual(items[0]["url"], "https://example.com/audio_0916.mp3")

    def test_sqlite_catalog_idempotency_workflow(self):
        """Catalog prevents duplicate transcriptions unless force=True."""
        url = "https://example.com/town_hearing.mp4"
        item_id = "hearing_01"

        # 1. Register new item
        registered = self.catalog.register_item(item_id, url, "Town Hearing 01", "channel_1")
        self.assertTrue(registered)

        # 2. Duplicate registration returns False
        dup = self.catalog.register_item(item_id, url, "Town Hearing 01", "channel_1")
        self.assertFalse(dup)

        # 3. Initially not transcribed
        self.assertFalse(self.catalog.is_transcribed(item_id))

        # 4. Mark completed
        self.catalog.mark_completed(item_id, "/tmp/transcripts/hearing_01")
        self.assertTrue(self.catalog.is_transcribed(item_id))

        # 5. Check stats
        stats = self.catalog.get_stats()
        self.assertEqual(stats.get("completed"), 1)

    def test_pipeline_v4_end_to_end_mock(self):
        """Pipeline v4 orchestrates ingestion, mocked ASR, FSM, and exports all 6 formats."""
        pipeline = TranscriptionPipelineV4(
            output_dir=self.tmp_dir / "transcripts",
            enable_streaming=False,
        )

        mock_segments = [
            {"start": 1.0, "end": 4.0, "ts": "00:00:01", "text": "I call this meeting to order with a roll call. Mr. MacLellan?"},
            {"start": 5.0, "end": 7.0, "ts": "00:00:05", "text": "Present."},
        ]

        # Create dummy source wav
        dummy_wav = self.tmp_dir / "dummy.wav"
        import wave
        with wave.open(str(dummy_wav), "wb") as w:
            w.setnchannels(1)
            w.setsampwidth(2)
            w.setframerate(16000)
            w.writeframes(b"\x00\x00" * 32000)

        with patch.object(pipeline.transcriber, "transcribe", return_value=mock_segments):
            res = pipeline.process(
                source=str(dummy_wav),
                playbook_name="municipal_meetings",
                job_id="test_job_e2e",
            )

        self.assertEqual(res["status"], "success")
        self.assertEqual(res["job_id"], "test_job_e2e")
        self.assertEqual(len(res["files"]), 6)
        self.assertTrue(Path(res["files"]["md"]).exists())
        self.assertTrue(Path(res["files"]["txt"]).exists())
        self.assertTrue(Path(res["files"]["srt"]).exists())
        self.assertTrue(Path(res["files"]["vtt"]).exists())
        self.assertTrue(Path(res["files"]["docx"]).exists())
        self.assertTrue(Path(res["files"]["json"]).exists())


    def test_rss_feed_malformed_xml_resilience(self):
        """Malformed XML with unescaped ampersands or control characters is sanitized and parsed."""
        malformed_rss = """<?xml version="1.0" encoding="UTF-8"?>
        <rss version="2.0">
          <channel>
            <title>Town & School Broadcast</title>
            <item>
              <title>Planning Board & Zoning \x02 Hearing</title>
              <guid>item_malformed_01</guid>
              <enclosure url="https://example.com/stream_01.mp3" type="audio/mpeg"/>
            </item>
          </channel>
        </rss>
        """
        items = self.ingestor.parse_rss_feed(malformed_rss)
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["id"], "item_malformed_01")
        self.assertIn("Planning Board & Zoning", items[0]["title"])
        self.assertEqual(items[0]["url"], "https://example.com/stream_01.mp3")

    def test_export_webvtt_escapes_markup_entities(self):
        """WebVTT voice span cues escape & and < in text to avoid invalid cue payloads."""
        blocks = [
            {"ts": "00:00:01", "start": 1.0, "end": 4.0, "speaker": "Speaker", "text": "Cost is < $500 & quality is > 90%."}
        ]
        meta = {"title": "Entity Test", "date": "2026-09-16", "source": "local"}
        paths = self.exporter.export(blocks, meta)
        content = paths["vtt"].read_text(encoding="utf-8")
        self.assertIn("&lt; $500 &amp; quality is &gt; 90%", content)
        self.assertNotIn("< $500", content)

    def test_discover_catalog_respects_limit(self):
        """discover_catalog slices discovered items to limit."""
        feed_xml = """<?xml version="1.0"?>
        <rss version="2.0"><channel>
          <item><title>Ep 1</title><guid>1</guid><enclosure url="http://a.com/1.mp3"/></item>
          <item><title>Ep 2</title><guid>2</guid><enclosure url="http://a.com/2.mp3"/></item>
          <item><title>Ep 3</title><guid>3</guid><enclosure url="http://a.com/3.mp3"/></item>
        </channel></rss>"""
        items = self.ingestor.discover_catalog(feed_xml, limit=2)
        self.assertEqual(len(items), 2)
        self.assertEqual(items[0]["id"], "1")
        self.assertEqual(items[1]["id"], "2")

    def test_extract_audio_retries_with_client_fallback_on_403(self):
        """When initial yt-dlp call yields 403, second strategy with player_client fallback succeeds."""
        job_id = "test_retry_job"
        job_dir = self.tmp_dir / job_id
        job_dir.mkdir(parents=True, exist_ok=True)

        call_count = {"count": 0}

        def fake_run(cmd, *args, **kwargs):
            if "yt-dlp" in cmd:
                call_count["count"] += 1
                if call_count["count"] == 1:
                    # First attempt fails with 403
                    raise subprocess.CalledProcessError(1, cmd, stderr=b"HTTP Error 403: Forbidden")
                else:
                    # Second attempt (fallback client) succeeds and writes opus stream
                    out_f = job_dir / "stream_raw.opus"
                    out_f.write_bytes(b"dummy opus data")
                    return subprocess.CompletedProcess(cmd, 0)
            elif "ffmpeg" in cmd:
                wav_target = Path(cmd[-1])
                wav_target.write_bytes(b"RIFF" + b"\x00" * 40)
                return subprocess.CompletedProcess(cmd, 0)
            return subprocess.CompletedProcess(cmd, 0)

        with patch.object(self.ingestor, "inspect_source", return_value={"title": "Test Stream"}):
            with patch("subprocess.run", side_effect=fake_run):
                wav_out, meta = self.ingestor.extract_audio("https://www.youtube.com/watch?v=0w7tT9V1KIA", job_id)

        self.assertEqual(call_count["count"], 2)
        self.assertTrue(wav_out.exists())
        self.assertEqual(wav_out.name, "audio_16k.wav")

    def test_extract_audio_resolves_arbitrary_stream_extensions(self):
        """extract_audio locates stream_raw regardless of whether yt-dlp extracts opus, webm, or m4a."""
        job_id = "test_ext_resolution"
        job_dir = self.tmp_dir / job_id
        job_dir.mkdir(parents=True, exist_ok=True)

        def fake_run(cmd, *args, **kwargs):
            if "yt-dlp" in cmd:
                out_f = job_dir / "stream_raw.opus"
                out_f.write_bytes(b"opus raw data")
                return subprocess.CompletedProcess(cmd, 0)
            elif "ffmpeg" in cmd:
                wav_target = Path(cmd[-1])
                wav_target.write_bytes(b"RIFF" + b"\x00" * 40)
                return subprocess.CompletedProcess(cmd, 0)
            return subprocess.CompletedProcess(cmd, 0)

        with patch.object(self.ingestor, "inspect_source", return_value={"title": "Opus Stream"}):
            with patch("subprocess.run", side_effect=fake_run):
                wav_out, meta = self.ingestor.extract_audio("https://www.youtube.com/watch?v=test", job_id)

        self.assertTrue(wav_out.exists())


if __name__ == "__main__":
    unittest.main()
